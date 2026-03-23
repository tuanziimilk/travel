#!/usr/bin/env python3
from __future__ import annotations

import argparse
import ast
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook


def normalize_country(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "us" if text.upper() == "MAIN" else text


def clean_brand(text: str) -> str:
    brand = re.sub(r"\s*\([^)]*\)", "", text).strip()
    brand = re.sub(r"\s*,\s*a\s+[^,]+$", "", brand, flags=re.I)
    brand = re.sub(r"\s+(generally|primarily)$", "", brand, flags=re.I)
    return brand.strip(" .,:;")


def lead_text(details: str) -> str:
    text = re.split(r"Key Details|Negative feedback", details, maxsplit=1, flags=re.I)[0]
    sentences = re.split(r"(?<=[.!?])\s+", text.strip())
    return " ".join(sentences[:2]).strip()


def merchant_name(term_name: object, domain: object, details: str = "") -> str:
    if term_name and str(term_name).strip():
        return clean_brand(str(term_name).strip())
    if details:
        head = lead_text(details)
        patterns = [
            r"^(?:Yes,\s*)?([A-Z0-9][A-Za-z0-9.&'() -]+?)\s+offers\s+free(?:[\s,]+\w+){0,5}\s+shipping",
            r"^(?:Based on [^,]+,\s*)?(?:No,\s*)?([A-Z0-9][A-Za-z0-9.&'() -]+?)\s+offers\s+free(?:[\s,]+\w+){0,5}\s+shipping",
            r"^(?:Yes,\s*)?([A-Z][A-Za-z0-9.&' -]+?)\s+offers\s+free(?:\s+\w+){0,3}\s+shipping",
            r"^(?:Based on [^,]+,\s*)?([A-Z0-9][A-Za-z0-9.&'() -]+?)\s+primarily deals with digital products",
            r"^(?:No,\s*)?([A-Z0-9][A-Za-z0-9.&'() -]+?)\s+does not(?:\s+\w+){0,6}\s+offer(?:\s+\w+){0,2}\s+free shipping",
            r"^(?:No,\s*)?([A-Z0-9][A-Za-z0-9.&'() -]+?)(?:,\s+[^,]+,\s+|\s+)does not(?:\s+\w+){0,6}\s+offer(?:\s+\w+){0,2}\s+free shipping",
            r"^(?:Based on [^,]+,\s*)?(?:No,\s*)?([A-Z0-9][A-Za-z0-9.&'() -]+?)(?:,\s+[^,]+)?\s+does not(?:\s+\w+){0,6}\s+(?:offer(?:\s+\w+){0,2}\s+free shipping|provide shipping|explicitly advertise)",
            r"^([A-Z][A-Za-z0-9.&' -]+?)\s+does not(?:\s+\w+){0,4}\s+offer free shipping",
            r"^([A-Z][A-Za-z0-9.&' -]+?)\s+generally does not offer free shipping",
            r"^([A-Z][A-Za-z0-9.&' -]+?)\s+does not generally advertise a sitewide free shipping policy",
            r"^(?:Based on [^,]+,\s*)?([A-Z0-9][A-Za-z0-9.&'() -]+?)\s+is(?:\s+\w+){0,6}\s+(?:digital|a proprietary|primarily a digital)",
            r"^([A-Z][A-Za-z0-9.&' -]+?)\s+primarily deals with digital products",
        ]
        for pattern in patterns:
            match = re.search(pattern, head)
            if match:
                return clean_brand(match.group(1))
    if not domain:
        return ""
    text = str(domain).strip()
    if text.startswith("www."):
        text = text[4:]
    base = text.split(".")[0]
    return clean_brand(" ".join(part.capitalize() for part in base.split("-")))


def get_field(row: dict[str, object], *names: str) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def extract_snippet(detail_text: str) -> str:
    if not detail_text:
        return ""
    try:
        parsed = ast.literal_eval(detail_text)
        if isinstance(parsed, list) and parsed:
            return str(parsed[0])
    except Exception:
        pass
    return detail_text


def extract_threshold(text: str) -> str:
    match = re.search(r"(?:over|orders? of|orders? above|minimum of|spend)\s+(\$?\d+(?:\.\d+)?)", text, re.I)
    if match:
        return match.group(1)
    return ""


def match_group(text: str, pattern: str) -> str:
    match = re.search(pattern, text, re.I)
    return match.group(1).strip() if match else ""


def clean_fragment(text: str) -> str:
    text = text.strip().strip(" .,:;")
    text = re.sub(r"^(while|and|but)\s+", "", text, flags=re.I)
    return text


def sentence_case(text: str) -> str:
    text = clean_fragment(text)
    if not text:
        return ""
    return text[0].upper() + text[1:]


def product_scope_detail(text: str) -> str:
    scope = match_group(text, r"including ([^.]+)")
    if scope:
        scope = clean_fragment(scope)
        if ", with typical delivery" in scope.lower():
            scope = re.split(r",\s+with\s+typical delivery", scope, maxsplit=1, flags=re.I)[0]
        return f"It mainly covers {scope}."
    eligible = match_group(text, r"Eligible items include ([^.]+)")
    if eligible:
        eligible = clean_fragment(eligible)
        if ", with typical delivery" in eligible.lower():
            eligible = re.split(r",\s+with\s+typical delivery", eligible, maxsplit=1, flags=re.I)[0]
        return f"Eligible items include {eligible}."
    scope = match_group(text, r"This applies to most products across their site, including ([^.]+)")
    if scope:
        return f"It mainly covers {clean_fragment(scope)}."
    return ""


def region_detail(text: str) -> str:
    lowered = text.lower()
    if "continental united states" in lowered:
        return "It applies within the continental U.S. only."
    if "contiguous united states" in lowered:
        if "alaska and hawaii" in lowered:
            return "It applies in the contiguous U.S. and excludes Alaska and Hawaii."
        return "It applies within the contiguous U.S. only."
    if "contiguous 48 states" in lowered:
        return "It applies within the contiguous 48 states only."
    countries = match_group(text, r"to the ([^.]+)")
    if countries and any(token in countries.lower() for token in ["u.s.", "canada", "australia", "u.k"]):
        return f"It is available to {clean_fragment(countries)}."
    return ""


def program_detail(text: str) -> str:
    value = match_group(text, r"Members of ([A-Z][A-Za-z0-9&' +.-]+) can receive free expedited delivery")
    if value:
        return f"{value} members can also get free expedited delivery."
    value = match_group(text, r"([A-Z][A-Za-z0-9&' +.-]+ members) receive free shipping with no minimum")
    if value:
        return f"{clean_fragment(value)} get free shipping with no minimum."
    value = match_group(text, r"(?:Recurring deliveries through (?:their )?)?([A-Z][A-Za-z0-9&' +.-]+ program) also qualif(?:y|ies) for free shipping")
    if value:
        return f"The {clean_fragment(value)} also qualifies for free shipping."
    value = match_group(text, r"through ([A-Z][A-Za-z0-9&' +.-]+ membership)")
    if value:
        return f"It is available through {clean_fragment(value)}."
    if "all subscription orders" in text.lower():
        return "Subscription orders also qualify for free shipping."
    return ""


def logistics_detail(text: str) -> str:
    lowered = text.lower()
    if "usps ground advantage shipping" in lowered:
        return "It uses USPS Ground Advantage shipping."
    if "in-store pickup" in lowered:
        return "It also offers in-store pickup for eligible orders."
    if "insured fedex shipping" in lowered:
        return "Orders ship with insured FedEx delivery."
    if "require a signature upon delivery" in lowered or "signature upon delivery" in lowered:
        return "Orders usually require a signature on delivery."
    if "free return shipping" in lowered:
        return "It also includes free return shipping on returns and exchanges."
    if "same-day delivery" in lowered:
        return "Walmart+ also includes free same-day delivery from local stores."
    if "2-day or next-day shipping" in lowered:
        return "Some items also qualify for free 2-day or next-day shipping."
    return ""


def digital_perk_detail(text: str) -> str:
    lowered = text.lower()
    if "3-day free trial" in lowered or "7-day free trial" in lowered:
        return "It offers a 3-day desktop trial and a 7-day mobile trial."
    if "titan business email" in lowered or "ssl certificates" in lowered:
        return "It instead includes free perks like Titan Business Email or SSL certificates."
    if "free, drm-free games" in lowered or "free games" in lowered:
        return "It frequently offers free DRM-free games through digital download."
    if "free registration" in lowered or "$0,00 for u.s. bank direct deposits" in lowered or "$0.00 for u.s. bank direct deposits" in lowered:
        return "Users typically pay platform fees instead of shipping charges."
    return ""


def exception_detail(text: str) -> str:
    lowered = text.lower()
    code = match_group(text, r'promo code "?([A-Z0-9]+)"?')
    if code:
        return f"It usually requires promo code {code}."
    code = match_group(text, r'code like "?([A-Z0-9]+)"?')
    if code:
        return f"It usually requires code {code}."
    if "temporary promotional codes" in lowered:
        return "Any free shipping usually depends on temporary promo codes."
    if "orders under" in lowered:
        fee = match_group(text, r"orders under [^.]*?costs? ([\$€£]\d+(?:\.\d+)?)")
        if fee:
            return f"Orders below that amount usually cost {fee} to ship."
        fee = match_group(text, r"orders under [^.]*?flat-rate shipping is ([\$€£]\d+(?:\.\d+)?)")
        if fee:
            return f"Orders below that amount usually add {fee} shipping."
    if "calculated at checkout" in lowered:
        return "Shipping costs are calculated at checkout on its official site."
    if "$0 delivery fees" in lowered and "uber eats" in lowered:
        return "Some locations may show $0 delivery fees through Uber Eats or Uber One."
    if "third-party retailers" in lowered:
        return "Free shipping is more likely through third-party retailers than direct orders."
    if "official airvape usa site" in lowered or "airvape usa site" in lowered:
        return "The official U.S. site usually charges shipping, while some retailers offer it free."
    if "vendor-direct" in lowered or "factory-direct" in lowered or "oversized" in lowered:
        return "Vendor-direct, factory-direct, or oversized items are usually excluded."
    if "u.s. territories" in lowered:
        return "It excludes U.S. territories where shipping is restricted."
    if "month-to-month subscriptions incur" in lowered:
        return "Month-to-month subscriptions do not qualify for free shipping."
    return ""


def trim_words(text: str, max_words: int = 20) -> str:
    words = text.split()
    if len(words) <= max_words:
        return text
    return " ".join(words[:max_words]).rstrip(".,;:") + "."


def has_no_minimum(text: str) -> bool:
    lowered = text.lower()
    return "no minimum purchase" in lowered or "no minimum order" in lowered or "all orders" in lowered


def extract_shipping_fee(text: str) -> str:
    match = re.search(r"(?:shipping (?:for orders .*? )?(?:generally costs|costs|fee(?:s)? (?:can be|is|are|of)?|charges apply(?: based on .*?)?)[^\$€£]*)([\$€£]\d+(?:\.\d+)?)", text, re.I)
    if match:
        return match.group(1)
    return ""


def extract_delivery_window(text: str) -> str:
    match = re.search(r"(\d+\s*(?:-|to)\s*\d+\s*(?:business|working)\s+days)", text, re.I)
    if match:
        return match.group(1)
    return ""


def first_sentence_positive(brand: str, text: str) -> str:
    lowered = text.lower()
    threshold = extract_threshold(text)

    if "prepaid multi-month subscription plans" in lowered:
        sentence = f"Yes, {brand} offers free shipping on prepaid multi-month subscription plans."
        if "contiguous united states" in lowered:
            sentence = f"Yes, {brand} offers free shipping on prepaid multi-month subscription plans in the contiguous U.S."
        return sentence

    if "all orders" in lowered and "free standard shipping" in lowered:
        return f"Yes, {brand} offers free standard shipping on all orders."

    if "all orders" in lowered and "free, insured fedex shipping" in lowered:
        return f"Yes, {brand} offers free insured FedEx shipping on all orders."

    if "all orders" in lowered and "free shipping" in lowered:
        return f"Yes, {brand} offers free shipping on all orders."

    if "free basic shipping" in lowered and threshold:
        code = match_group(text, r'like "?([A-Z0-9]+)"?')
        if code:
            return f"Yes, {brand} offers free basic shipping on orders over {threshold} with code {code}."
        return f"Yes, {brand} offers free basic shipping on orders over {threshold}."

    if threshold and "free standard ground shipping" in lowered:
        if "continental u.s" in lowered or "continental united states" in lowered:
            return f"Yes, {brand} offers free standard ground shipping on orders over {threshold} in the continental U.S."
        return f"Yes, {brand} offers free standard ground shipping on orders over {threshold}."

    if threshold and "free standard shipping" in lowered:
        if "continental united states" in lowered:
            return f"Yes, {brand} offers free standard shipping on orders over {threshold} in the continental U.S."
        return f"Yes, {brand} offers free standard shipping on orders over {threshold}."

    if threshold and "free ground shipping" in lowered:
        if "promo code" in lowered:
            code = match_group(text, r'promo code "?([A-Z0-9]+)"?')
            if code:
                return f"Yes, {brand} offers free ground shipping on orders over {threshold} with promo code {code}."
        if "contiguous 48 states" in lowered:
            return f"Yes, {brand} offers free ground shipping on orders over {threshold} in the contiguous 48 states."
        if "contiguous united states" in lowered:
            return f"Yes, {brand} offers free ground shipping on orders over {threshold} in the contiguous U.S."
        return f"Yes, {brand} offers free ground shipping on orders over {threshold}."

    if threshold and "free shipping" in lowered:
        if "basic shipping" in lowered:
            code = match_group(text, r'like "?([A-Z0-9]+)"?')
            if code:
                return f"Yes, {brand} offers free basic shipping on orders over {threshold} with code {code}."
            return f"Yes, {brand} offers free basic shipping on orders over {threshold}."
        if "contiguous united states" in lowered:
            return f"Yes, {brand} offers free shipping on orders over {threshold} in the contiguous U.S."
        return f"Yes, {brand} offers free shipping on orders over {threshold}."

    if "depends on the region and minimum order value" in lowered or "official airvape usa site" in lowered:
        return f"{brand} does not offer sitewide free shipping."

    if "hardware devices and bundles" in lowered:
        return f"Yes, {brand} offers free shipping on hardware devices and bundles bought directly from its website."

    if "select items" in lowered or "qualifying orders" in lowered:
        return f"Yes, {brand} offers free shipping on select items or qualifying orders."

    if "free shipping" in lowered and "standard shipping" in lowered:
        return f"Yes, {brand} offers free standard shipping."

    if "free shipping" in lowered:
        return f"Yes, {brand} offers free shipping."

    return f"Yes, {brand} offers free shipping under certain conditions."


def specific_exception_detail(text: str) -> str:
    lowered = text.lower()
    if "vendor-direct" in lowered or "factory-direct" in lowered or "oversized" in lowered:
        return "Vendor-direct, factory-direct, or oversized items are usually excluded."
    fee = extract_shipping_fee(text)
    if "orders below" in lowered and fee:
        return f"Orders below that amount usually cost {fee} to ship."
    if "month-to-month subscriptions incur" in lowered:
        fee2 = match_group(text, r"month-to-month subscriptions incur a ([\$€£]\d+(?:\.\d+)?) shipping fee")
        if fee2:
            return f"Month-to-month subscriptions usually add {fee2} shipping."
    if "non-contiguous us states or canada" in lowered:
        return "Shipments to non-contiguous U.S. states or Canada cost extra."
    return ""


def detail_threshold(text: str) -> str:
    lowered = text.lower()
    for candidate in [
        program_detail(text),
        product_scope_detail(text),
        region_detail(text),
        logistics_detail(text),
        specific_exception_detail(text),
    ]:
        if candidate:
            return candidate
    if "free standard ground shipping" in lowered:
        window = extract_delivery_window(text)
        if window:
            return f"Eligible items usually arrive within {window}."
        return "It usually covers standard ground shipping."
    if "free ground shipping" in lowered:
        region = match_group(text, r"apply to (the [^.]+?states)")
        if region:
            return f"It usually covers ground shipping in {region}."
        return "It usually applies to ground shipping."
    if "free standard shipping" in lowered:
        if has_no_minimum(text):
            return "No minimum purchase is required."
        window = extract_delivery_window(text)
        if window:
            return f"Standard delivery usually takes {window}."
        return "It usually covers standard shipping."
    if "heat wave eyewear" in lowered:
        return "Some Heat Wave Eyewear items also ship free."
    if "select tours and travel packages" in lowered:
        return "It mainly applies to select tours and travel packages."
    if "subtotal after coupons" in lowered:
        threshold = extract_threshold(text)
        if threshold:
            return f"The {threshold} minimum is based on the subtotal after coupons."
        return "The minimum is based on the subtotal after coupons."
    return pick_variant(
        text,
        [
            "It applies to qualifying orders only.",
            "The offer usually covers eligible purchases.",
            "This offer is typically limited to qualifying orders.",
        ],
    )


def detail_membership(text: str, brand: str) -> str:
    program = match_group(text, r"members of (?:their )?([A-Z][A-Za-z0-9&' -]+?)\s+(?:subscription|membership)")
    if program:
        return f"That perk is tied to its {program} program."
    extra = program_detail(text)
    if extra:
        return extra
    return pick_variant(
        brand,
        [
            f"{brand} only includes free shipping through its membership or subscription program.",
            f"Free shipping is limited to {brand}'s membership or subscription program.",
            f"The free shipping benefit is mainly tied to {brand}'s membership offering.",
        ],
    )


def detail_merchant_dependent(text: str) -> str:
    if "local restaurants, alcohol, and laundry" in text.lower():
        return "Fees and delivery offers vary by the local restaurant or store."
    extra = exception_detail(text)
    if extra:
        return extra
    return pick_variant(
        "merchant-dependent",
        [
            "Free shipping or delivery depends on the individual merchant or store.",
            "Shipping terms vary by seller, so free shipping is not a general sitewide policy.",
            "Any free delivery offer usually comes from the individual store rather than the platform as a whole.",
        ],
    )


def detail_digital(text: str) -> str:
    lowered = text.lower()
    extra = digital_perk_detail(text)
    if extra:
        return extra
    if "booking confirmation" in lowered or "booking confirmations" in lowered:
        return "Bookings are confirmed digitally instead of being shipped."
    if "digital downloads" in lowered or "download directly" in lowered:
        return "Its products are usually delivered through digital downloads."
    if "domains and hosting" in lowered:
        return "Its core services are domains and hosting rather than physical goods."
    return pick_variant(
        text,
        [
            "It provides digital products or services instead.",
            "Its products or confirmations are delivered digitally rather than shipped.",
            "Customers typically receive downloads, bookings, or confirmations instead of shipped items.",
        ],
    )


def detail_explicit_no(text: str) -> str:
    lowered = text.lower()
    if "vehicle rental service" in lowered:
        return "It is a vehicle rental service rather than a product retailer."
    extra = exception_detail(text)
    if extra:
        return extra
    fee = extract_shipping_fee(text)
    if fee:
        return f"Standard shipping generally costs {fee}."
    return pick_variant(
        text,
        [
            "Paid shipping options apply instead.",
            "Standard delivery charges apply instead of a free shipping offer.",
            "Orders generally use paid shipping rather than a free shipping policy.",
        ],
    )


def detail_positive_general(text: str) -> str:
    lowered = text.lower()
    for candidate in [
        specific_exception_detail(text),
        logistics_detail(text),
        program_detail(text),
        product_scope_detail(text),
        region_detail(text),
        exception_detail(text),
    ]:
        if candidate:
            return candidate
    if "30-day returns" in lowered:
        return "It also offers 30-day returns on eligible orders."
    if has_no_minimum(text):
        if "standard ground shipping" in lowered:
            return "Standard ground shipping is available with no minimum purchase."
        if "standard shipping" in lowered:
            return "Standard shipping is available with no minimum purchase."
        if "ground shipping" in lowered:
            return "Ground shipping is available with no minimum purchase."
        return "No minimum purchase is required."
    region = match_group(text, r"apply to (the [^.]+?states)")
    if region:
        return f"The offer is mainly available in {region}."
    if "ground shipping" in lowered:
        return "It usually applies to ground shipping."
    if "standard shipping" in lowered:
        return "It typically applies to standard shipping."
    return pick_variant(
        text,
        [
            "It usually applies to standard or ground shipping only.",
            "The offer typically applies under limited shipping conditions.",
            "Free shipping is usually tied to specific shipping methods or eligible orders.",
        ],
    )


def detail_select_items(text: str) -> str:
    lowered = text.lower()
    threshold = extract_threshold(text)
    if "orders over" in lowered and threshold and "ground advantage" in lowered:
        return "Orders over that amount usually ship with USPS Ground Advantage."
    if "travel packages" in lowered:
        return "It is mostly available on select tours and travel packages."
    scope = product_scope_detail(text)
    if scope:
        return scope
    return pick_variant(
        text,
        [
            "Scope restrictions may apply.",
            "The offer is usually limited to eligible items or orders.",
            "Availability can vary by item and order type.",
        ],
    )


def pick_variant(brand: str, options: list[str]) -> str:
    if not options:
        return ""
    return options[sum(ord(char) for char in brand) % len(options)]


def finalize_answer(first_sentence: str, second_sentence: str = "") -> str:
    if not second_sentence:
        return first_sentence
    combined = f"{first_sentence} {second_sentence}"
    if len(combined.split()) <= 50:
        return combined
    return f"{first_sentence} {trim_words(second_sentence, 50 - len(first_sentence.split()))}"


def finalize_sentences(*sentences: str) -> str:
    parts = [clean_fragment(sentence) + "." for sentence in sentences if sentence and clean_fragment(sentence)]
    if not parts:
        return ""
    answer = parts[0]
    for part in parts[1:]:
        candidate = f"{answer} {part}"
        if len(candidate.split()) <= 50:
            answer = candidate
    return answer


def format_yes_threshold(brand: str, threshold: str, details: str) -> str:
    return finalize_answer(first_sentence_positive(brand, details), detail_threshold(details))


def format_limited_membership(brand: str, details: str) -> str:
    return finalize_answer(f"{brand} does not offer sitewide free shipping.", detail_membership(details, brand))


def format_limited_merchant_dependent(brand: str, details: str) -> str:
    return finalize_answer(f"{brand} does not offer sitewide free shipping.", detail_merchant_dependent(details))


def format_digital_service(brand: str, details: str) -> str:
    lowered = details.lower()
    opening_options = [
        f"No, {brand} does not offer free shipping because it does not ship physical products.",
        f"{brand} is a digital service, so free shipping does not apply.",
        f"{brand} is not a physical retailer, so free shipping does not apply.",
    ]
    if "digital software service" in lowered or "software-based solution" in lowered:
        base = f"{brand} is a digital software service, not a physical retailer, so free shipping does not apply."
    elif "freelancing platform" in lowered or "digital services" in lowered:
        base = f"{brand} is a digital services platform, so free shipping does not apply."
    elif "digital distribution platform" in lowered:
        base = f"{brand} is a digital download platform, so free shipping does not apply."
    elif "domains and hosting" in lowered:
        base = f"{brand} provides digital services rather than physical products, so free shipping does not apply."
    else:
        base = pick_variant(brand, opening_options)
    extra = detail_digital(details)
    third = ""
    if "domains and hosting" in details.lower():
        third = "Products are delivered digitally instead of being shipped."
    elif "digital services" in details.lower() or "freelancing platform" in details.lower():
        third = "Projects and services are delivered through the platform."
    return finalize_sentences(
        base,
        extra,
        third,
    )


def format_explicit_no(brand: str, details: str) -> str:
    return finalize_answer(f"No, {brand} does not offer free shipping.", detail_explicit_no(details))


def format_positive_general(brand: str, details: str) -> str:
    return finalize_answer(first_sentence_positive(brand, details), detail_positive_general(details))


def format_positive_no_minimum(brand: str, details: str) -> str:
    second = detail_positive_general(details)
    if second == "Standard shipping is available with no minimum purchase.":
        second = "It usually covers standard shipping."
    if second == "Ground shipping is available with no minimum purchase.":
        second = "It usually covers ground shipping."
    if second == "Standard ground shipping is available with no minimum purchase.":
        second = "It usually covers standard ground shipping."
    return finalize_answer(first_sentence_positive(brand, details), second)


def format_positive_select_items(brand: str, details: str) -> str:
    return finalize_answer(
        f"Yes, {brand} offers free shipping on select items or qualifying orders.",
        detail_select_items(details),
    )


def detail_unclear(text: str) -> str:
    for candidate in [
        exception_detail(text),
        product_scope_detail(text),
        region_detail(text),
        logistics_detail(text),
    ]:
        if candidate:
            return candidate
    return ""


def format_unclear(brand: str, details: str) -> str:
    return finalize_answer(f"{brand} does not clearly advertise free shipping.", detail_unclear(details))


def classify_answer(brand: str, details: str) -> str:
    head = lead_text(details)
    text = head.lower()
    threshold = extract_threshold(head)
    no_minimum = has_no_minimum(head)
    positive_phrase = bool(re.search(r"offers free(?:[\s,]+\w+){0,5}\s+shipping", text))
    explicit_negative = any(
        phrase in text
        for phrase in [
            "does not typically offer free shipping",
            "does not offer free shipping",
            "does not offer universal free shipping",
            "does not explicitly advertise",
            "free shipping is not generally offered",
        ]
    )

    if "depends on the region and minimum order value" in text or "official airvape usa site" in details.lower():
        return finalize_answer(f"{brand} does not offer sitewide free shipping.", exception_detail(details))

    if threshold and positive_phrase:
        return format_yes_threshold(brand, threshold, details)

    if no_minimum and positive_phrase:
        return format_positive_no_minimum(brand, details)

    if any(phrase in text for phrase in ["does not ship physical products", "digital platform", "booking accommodation", "booking confirmations", "provided through downloads", "delivers software digitally", "digital software company", "digital downloads"]):
        return format_digital_service(brand, details)
    if any(
        phrase in text
        for phrase in [
            "digital software service",
            "digital distribution platform",
            "digital products like domains and hosting",
            "digital services",
            "freelancing platform",
            "not a retailer of physical goods",
            "does not involve physical products",
            "no physical product",
            "there is no shipping involved",
        ]
    ):
        return format_digital_service(brand, details)

    if "depends on the individual merchant" in text or "depends on the specific local store" in text or "entirely dependent on the specific local store" in text:
        return format_limited_merchant_dependent(brand, details)

    if "does not offer free shipping as a standard" in text or "does not generally advertise a sitewide free shipping policy" in text:
        if "membership" in text or "subscription" in text or "members of their" in text:
            return format_limited_membership(brand, details)
        return format_limited_merchant_dependent(brand, details)

    if "free shipping exclusively for members" in text or "members of their" in text:
        return format_limited_membership(brand, details)

    if "offers free shipping on select items" in text or "qualifying orders" in text:
        return format_positive_select_items(brand, details)

    if explicit_negative:
        return format_explicit_no(brand, details)

    if positive_phrase:
        return format_positive_general(brand, details)

    return format_unclear(brand, details)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate free shipping FAQ output rows from an upload template.")
    parser.add_argument("input", type=Path, help="Path to upload-template.xlsx")
    parser.add_argument("output", type=Path, help="Path to output xlsx")
    args = parser.parse_args()

    wb = load_workbook(args.input)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("Input workbook is empty.")

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    data_rows = [dict(zip(headers, row)) for row in rows[1:] if any(cell is not None for cell in row)]

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Sheet1"
    out_ws.append([
        "Content",
        "Country",
        "TermId",
        "TermName",
        "Domain",
        "Source",
        "Subclo",
        "板块名称",
        "Title1",
        "Brief Introduction",
        "Href Kw",
        "Href Url",
    ])

    for row in data_rows:
        details = extract_snippet(get_field(row, "discount_details"))
        brand = merchant_name(row.get("term_name"), row.get("domain"), details)
        country = normalize_country(row.get("country"))
        fact_type = get_field(row, "fact_type")
        intro = classify_answer(brand, details)
        out_ws.append([
            "faq",
            country,
            get_field(row, "term_id"),
            brand,
            get_field(row, "domain"),
            "Ai",
            fact_type,
            "faq",
            f"Does {brand} offer free shipping?",
            intro,
            "",
            "",
        ])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(args.output)


if __name__ == "__main__":
    main()
